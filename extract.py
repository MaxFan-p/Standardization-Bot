"""Download files from Slack and extract text from Word documents."""
import io
import urllib.request

import config


def download_slack_file(url_private_download: str) -> bytes:
    """Fetch a Slack-hosted file. Requires the bot token in the auth header."""
    req = urllib.request.Request(
        url_private_download,
        headers={"Authorization": f"Bearer {config.SLACK_BOT_TOKEN}"},
    )
    with urllib.request.urlopen(req) as resp:
        return resp.read()


def extract_docx_text(data: bytes) -> str:
    """Pull text (paragraphs + table cells) out of a .docx file."""
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)